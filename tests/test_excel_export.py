import copy
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from excel_export import export_xlsx


def fixture():
    return {'source_name':'=literal-report.PDF','sample_label':'=literal-report','units':'uCi/Unit','sigma':'1.000',
            'generated':'9/16/2026 11:42:11 AM','acquired':'6/30/2026 11:25:16 AM','sha256':'a'*64,
            'rows':[
                {'nuclide':'Pt-189','confidence':'0.977','activity':'2.0022751E+03','uncertainty':'2.1220114E+02',
                 'flags':'@','status':'Reported','page':36,'raw_line':'Pt-189 @ 0.977 2.0022751E+03 2.1220114E+02'},
                {'nuclide':'Os-183','confidence':'0.931','activity':None,'uncertainty':None,
                 'flags':'X','status':'Rejected by interference analysis','page':36,'raw_line':'X Os-183 0.931'},
                {'nuclide':'Ir-185','confidence':'0.700','activity':'1.0E+01','uncertainty':'1.0E+00',
                 'flags':'','status':'Reported','page':36,'raw_line':'Ir-185 0.700 1.0E+01 1.0E+00'}],
            'line_rows':[{'nuclide':'Pt-189','energy':'721.22','yield_percent':'17.20','activity':None,'uncertainty':None,
                          'status':'Not found in spectrum','flags':'','page':34,'raw_line':'721.22 not found'}]}


class ExcelExportTests(unittest.TestCase):
    def test_native_workbook_preserves_values_missing_data_and_literal_text(self):
        report=fixture()
        original=copy.deepcopy(report)
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'result.xlsx'
            self.assertEqual(export_xlsx([report],path),2)
            book=load_workbook(path)
            self.assertEqual(book.sheetnames,['Results','Source records','Energy lines'])
            result,source,lines=(book[s] for s in book.sheetnames)
            self.assertEqual(result['A8'].data_type,'s')
            self.assertEqual(result['A8'].value,'=literal-report')
            self.assertEqual(result['C8'].value,2002.2751)
            self.assertIsNone(result['C9'].value)
            self.assertIsNone(result['D9'].value)
            self.assertEqual(source['E2'].value,'2.0022751E+03')
            self.assertEqual(source['E2'].data_type,'s')
            self.assertEqual(source['J2'].value,'@')
            self.assertEqual(source['P4'].value,'Filtered: confidence <= 0.7')
            self.assertEqual(source['L2'].value.year,2026)
            self.assertEqual(source['L2'].value.hour,11)
            self.assertEqual(result.freeze_panes,'C8')
            self.assertEqual(len(result.tables),1)
            self.assertIn('>$I$4',result['G8'].value)
            self.assertEqual(result['I4'].value,0.1)
            self.assertIsNone(lines['E2'].value)
            self.assertEqual(lines['D2'].value,0.172)
            book.close()
        self.assertEqual(report,original)

    def test_all_confidences_and_custom_uncertainty_cutoff(self):
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'result.xlsx'
            self.assertEqual(export_xlsx([fixture()],path,include_below=True,uncertainty_threshold='15'),3)
            book=load_workbook(path)
            self.assertEqual(book['Results']['I4'].value,0.15)
            self.assertEqual(book['Source records']['P4'].value,'Included')
            book.close()

    def test_mixed_units_and_unrepresentable_values_are_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            target=Path(temp)/'result.xlsx'
            other=fixture();other['units']='Bq'
            with self.assertRaisesRegex(ValueError,'different activity units'):
                export_xlsx([fixture(),other],target)
            report=fixture();report['rows'][0]['activity']='1.2345678901234567'
            with self.assertRaisesRegex(ValueError,'precision'):
                export_xlsx([report],target)
            self.assertFalse(target.exists())
