"""Native Excel export. Reported values are never recalculated or decay-corrected."""
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from core import selected, review_threshold, atomic_write

NAVY = '253C59'
SCIENCE = '0.0000000E+00'


def number(value):
    if value is None or value == '':
        return None
    parsed = Decimal(str(value))
    if not parsed.is_finite() or len(parsed.normalize().as_tuple().digits) > 15:
        raise ValueError('A value exceeds Excel numeric precision. Use the exact source records or CSV.')
    converted = float(parsed)
    if Decimal(str(converted)) != parsed:
        raise ValueError('A value cannot be represented safely in Excel. Use the exact source records or CSV.')
    return converted


def stamp(value):
    if not value:
        return None
    try:
        return datetime.strptime(value, '%m/%d/%Y %I:%M:%S %p')
    except ValueError:
        return value  # Preserve an unfamiliar date rather than guessing its meaning.


def text(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'  # Literal text, including GENIE @ flags and formula-looking filenames.


def write_row(sheet, index, values):
    for col, value in enumerate(values, 1):
        text(sheet.cell(index, col), value)


def format_table(sheet, header_row, last_row, widths, name, frozen):
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = frozen
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    for row in sheet.iter_rows(min_row=1, max_row=last_row, max_col=len(widths)):
        for cell in row:
            cell.font = Font(name='Arial', size=10, color='233142')
            cell.alignment = Alignment(vertical='center')
    for row in range(header_row + 1, last_row + 1):
        sheet.row_dimensions[row].height = 24
    for cell in sheet[header_row]:
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    sheet.row_dimensions[header_row].height = 38
    # A header-only worksheet is valid when there are no selected results or energy lines.
    if last_row > header_row:
        table = Table(displayName=name, ref=f'A{header_row}:{get_column_letter(len(widths))}{last_row}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
        sheet.add_table(table)
    else:
        sheet.auto_filter.ref = f'A{header_row}:{get_column_letter(len(widths))}{header_row}'


def export_xlsx(reports, target, threshold='0.7', include_below=False, uncertainty_threshold='10'):
    if not reports:
        raise ValueError('Select at least one report.')
    selected([], threshold, include_below)
    cutoff = review_threshold(uncertainty_threshold)
    units = {r['units'] for r in reports}
    if len(units) != 1:
        raise ValueError('Reports have different activity units. Export each unit group separately.')
    unit = next(iter(units))
    sigmas = {str(r['sigma']) for r in reports}
    sigma_note = f", {next(iter(sigmas))} sigma" if len(sigmas) == 1 else ''
    workbook = Workbook()
    result = workbook.active
    result.title = 'Results'
    source = workbook.create_sheet('Source records')
    lines = workbook.create_sheet('Energy lines')
    result.sheet_properties.tabColor = NAVY
    write_row(result, 2, ['Platinum-189 weighted activities'])
    write_row(result, 4, ['Selection: all confidence values' if include_below else f'Selection: ID confidence > {threshold}'])
    text(result['H4'], 'Review above')
    result['I4'] = float(cutoff / 100)
    write_row(result, 5, ['Blank activity = unavailable. Original values, flags and dates are retained in Source records.'])
    write_row(result, 7, ['Report / sample', 'Nuclide', f'Activity ({unit})', f'Uncertainty ({unit}{sigma_note})',
                         'Relative uncertainty', 'ID confidence', 'Review', 'GENIE flags', 'Source'])
    write_row(source, 1, ['Report ref', 'Source PDF', 'Nuclide', 'Confidence (original text)', 'Activity (original text)',
                         'Uncertainty (original text)', 'Units', 'Sigma', 'GENIE status', 'GENIE flags', 'PDF page',
                         'Report generated', 'Acquisition started', 'Source SHA256', 'Original source row', 'Export selection'])
    write_row(lines, 1, ['Report ref', 'Nuclide', 'Energy (keV)', 'Yield (%)', f'Activity ({unit})',
                        f'Uncertainty ({unit})', 'GENIE line use', 'GENIE flags', 'PDF page', 'Original source row'])
    rr, sr, lr = 8, 2, 2
    for index, report in enumerate(reports, 1):
        ref = f'R{index:02d}'
        chosen = selected(report['rows'], threshold, include_below)
        for row in report['rows']:
            write_row(source, sr, [ref, report['source_name'], row['nuclide'], row['confidence'], row.get('activity'),
                                  row.get('uncertainty'), report['units'], str(report['sigma']), row['status'], row['flags'],
                                  row['page'], stamp(report['generated']), stamp(report['acquired']), report['sha256'],
                                  row.get('raw_line', ''), 'Included' if row in chosen else f'Filtered: confidence <= {threshold}'])
            for col in ['D', 'E', 'F', 'H']:
                source[f'{col}{sr}'].number_format = '@'
            for col in ['L', 'M']:
                source[f'{col}{sr}'].number_format = 'mm/dd/yyyy hh:mm:ss AM/PM'
            sr += 1
            if row not in chosen:
                continue
            write_row(result, rr, [report['sample_label'], row['nuclide'], number(row.get('activity')), number(row.get('uncertainty')),
                                  None, number(row['confidence']), None, row['flags'], f'{ref} / p{row["page"]}'])
            unavailable = 'Rejected' if row['status'].startswith('Rejected') else row['status']
            unavailable = unavailable.replace('"', '""')
            result[f'E{rr}'] = f'=IF(OR(C{rr}="",D{rr}="",C{rr}=0),"",ABS(D{rr}/C{rr}))'
            result[f'G{rr}'] = f'=IF(C{rr}="","{unavailable}",IF(E{rr}="","Cannot assess",IF(E{rr}>$I$4,"High uncertainty","Within threshold")))'
            for col in ['C', 'D']:
                result[f'{col}{rr}'].number_format = SCIENCE
            result[f'E{rr}'].number_format = '0.00%'
            result[f'F{rr}'].number_format = '0.000'
            rr += 1
        for row in report.get('line_rows', []):
            value = number(row.get('yield_percent'))
            write_row(lines, lr, [ref, row['nuclide'], number(row['energy']), value / 100 if value is not None else None,
                                 number(row.get('activity')), number(row.get('uncertainty')), row['status'], row['flags'],
                                 row['page'], row.get('raw_line', '')])
            lines[f'C{lr}'].number_format = '0.00'
            lines[f'D{lr}'].number_format = '0.00%'
            for col in ['E', 'F']:
                lines[f'{col}{lr}'].number_format = '0.000000E+00'
            lr += 1
    format_table(result, 7, rr - 1, [37, 12, 21, 23, 18, 15, 24, 12, 15], 'WeightedActivities', 'C8')
    format_table(source, 1, sr - 1, [10, 40, 12, 19, 23, 23, 14, 12, 35, 12, 10, 25, 25, 68, 78, 33], 'OriginalSummaryRecords', 'D2')
    format_table(lines, 1, lr - 1, [10, 12, 17, 15, 21, 21, 37, 12, 10, 95], 'OriginalEnergyLines', 'C2')
    result['A2'].font = Font(name='Arial', size=16, bold=True, color=NAVY)
    result['A5'].font = Font(name='Arial', size=10, italic=True, color='65758A')
    for cell in result[3][:9]:
        cell.border = Border(bottom=Side(style='thin', color=NAVY))
    result['I4'].number_format = '0.0%'
    result['I4'].fill = PatternFill('solid', fgColor='FFF1CB')
    if rr > 8:
        result.conditional_formatting.add(f'C8:G{rr-1}', FormulaRule(formula=['$G8="High uncertainty"'], fill=PatternFill('solid', fgColor='FFF1CB')))
        result.conditional_formatting.add(f'C8:G{rr-1}', FormulaRule(formula=['$G8="Rejected"'], fill=PatternFill('solid', fgColor='F8E7E6'), font=Font(color='8D3530')))
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    buffer = BytesIO()
    workbook.save(buffer)
    atomic_write(target, buffer.getvalue())
    return rr - 8
